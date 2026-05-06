import sys
import os
import re
import xml.etree.ElementTree as ET
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pathlib import Path
import copy
import matplotlib.colors as mcolors

ADAPTER_SCRIPT_PATH = Path(__file__).resolve()
SCRIPTS_DIR = ADAPTER_SCRIPT_PATH.parent
BEAMPY_SLIDES_ROOT = SCRIPTS_DIR.parent

if str(SCRIPTS_DIR) not in sys.path: sys.path.insert(0, str(SCRIPTS_DIR))

try:
    from svg_to_pptx.pptx_builder import create_pptx_with_native_svg
    HAS_BUILDER = True
except ImportError:
    HAS_BUILDER = False

def multiply_matrix(m1, m2):
    return [
        m1[0]*m2[0] + m1[2]*m2[1],
        m1[1]*m2[0] + m1[3]*m2[1],
        m1[0]*m2[2] + m1[2]*m2[3],
        m1[1]*m2[2] + m1[3]*m2[3],
        m1[0]*m2[4] + m1[2]*m2[5] + m1[4],
        m1[1]*m2[4] + m1[3]*m2[5] + m1[5]
    ]

def parse_transform_to_matrix(transform_str):
    m = [1.0, 0.0, 0.0, 1.0, 0.0, 0.0]
    if not transform_str: return m
    commands = re.findall(r'[a-zA-Z]+\([^)]+\)', transform_str)
    for cmd in commands:
        vals = [float(x) for x in re.findall(r'-?\d*\.?\d+', cmd)]
        if cmd.startswith("translate"):
            tx = vals[0]
            ty = vals[1] if len(vals) > 1 else 0
            m = multiply_matrix(m, [1, 0, 0, 1, tx, ty])
        elif cmd.startswith("scale"):
            sx = vals[0]
            sy = vals[1] if len(vals) > 1 else sx
            m = multiply_matrix(m, [sx, 0, 0, sy, 0, 0])
        elif cmd.startswith("matrix"):
            m = multiply_matrix(m, vals)
    return m

def absolute_flatten(root, plain_text_ids, target_w=1280, target_h=720, doc_w=800, doc_h=600):
    SVG_NS = "http://www.w3.org/2000/svg"
    ns_prefix = f"{{{SVG_NS}}}"
    
    defs_map = {}
    for defs in root.findall(f".//{ns_prefix}defs"):
        for child in defs:
            eid = child.get("id")
            if eid: defs_map[eid] = copy.deepcopy(child)

    scale = min(target_w / doc_w, target_h / doc_h)
    offset_x = (target_w - doc_w * scale) / 2
    offset_y = (target_h - doc_h * scale) / 2
    global_matrix = [scale, 0, 0, scale, offset_x, offset_y]

    new_root = ET.Element(f"{ns_prefix}svg", root.attrib)
    new_defs = ET.SubElement(new_root, f"{ns_prefix}defs")
    for def_elem in defs_map.values():
        new_defs.append(copy.deepcopy(def_elem))

    def walk(node, current_matrix, styles, current_parent):
        node_id = node.get("id")
        if node_id in plain_text_ids:
            return

        node_trans_str = node.get("transform", "")
        node_matrix = parse_transform_to_matrix(node_trans_str)
        cumulative_matrix = multiply_matrix(current_matrix, node_matrix)
        
        new_styles = styles.copy()
        for k, v in node.items():
            if k not in ("id", "transform", "class", "data-layer") and not k.startswith("{"):
                new_styles[k] = v
                
        tag = node.tag.replace(ns_prefix, "")
        is_module = tag == "g" and node_id and not (node_id.startswith("slide_") or node_id.startswith("layer_"))
        
        if is_module:
            new_group = ET.SubElement(current_parent, f"{ns_prefix}g")
            new_group.set("id", node_id)
            for child in node:
                if child.tag.endswith("defs"): continue
                walk(child, cumulative_matrix, new_styles, new_group)
            return

        if tag == "use":
            href = node.get("{http://www.w3.org/1999/xlink}href") or node.get("href")
            if href and href.startswith("#"):
                ref_id = href[1:]
                if ref_id in defs_map:
                    ux = float(node.get("x", "0"))
                    uy = float(node.get("y", "0"))
                    use_matrix = multiply_matrix(cumulative_matrix, [1, 0, 0, 1, ux, uy])
                    walk(defs_map[ref_id], use_matrix, new_styles, current_parent)
                    
        elif tag in ("path", "rect", "circle"):
            new_node = copy.deepcopy(node)
            if "transform" in new_node.attrib: del new_node.attrib["transform"]
            sx, sy, tx, ty = cumulative_matrix[0], cumulative_matrix[3], cumulative_matrix[4], cumulative_matrix[5]
            wrapper_g = ET.SubElement(current_parent, f"{ns_prefix}g")
            wrapper_g.set("transform", f"translate({tx:.4f},{ty:.4f}) scale({sx:.4f},{sy:.4f})")
            for k, v in new_styles.items():
                if k not in new_node.attrib: new_node.set(k, v)
            wrapper_g.append(new_node)
            
        elif tag in ("g", "svg"):
            for child in node:
                if child.tag.endswith("defs"): continue
                walk(child, cumulative_matrix, new_styles, current_parent)

    walk(root, global_matrix, {}, new_root)
    return new_root

def is_math(text_str):
    if not text_str: return False
    # Smart check: if it contains $ or complex math commands, it's vector math.
    # Otherwise, it's plain text (even if it has \& or \textbf)
    if "$" in text_str: return True
    math_indicators = [r'\frac', r'\sum', r'\int', r'\sqrt', r'^', r'_', r'\alpha', r'\beta']
    return any(cmd in text_str for cmd in math_indicators)

def clean_latex_for_pptx(text_str):
    if not text_str: return ""
    # 1. Unescape common characters
    res = text_str.replace(r'\&', '&').replace(r'\%', '%').replace(r'\#', '#').replace(r'\_', '_').replace(r'\\', ' ')
    # 2. Strip basic formatting commands like \textbf{...} -> ...
    res = re.sub(r'\\textbf\{([^}]*)\}', r'\1', res)
    res = re.sub(r'\\textit\{([^}]*)\}', r'\1', res)
    res = re.sub(r'\\text\{([^}]*)\}', r'\1', res)
    return res

def export_beampy_to_editable_pptx(output_pptx="presentation.pptx"):
    from beampy.document import document as bp_doc
    from beampy.exports import reset_module_rendered_flag, render_texts, save_layout
    
    debug_dir = Path("./beampy_debug_svgs")
    debug_dir.mkdir(exist_ok=True)
    
    reset_module_rendered_flag()
    render_texts()
    save_layout()
    
    doc_w = getattr(bp_doc, 'width', getattr(bp_doc, '_width', 800))
    doc_h = getattr(bp_doc, 'height', getattr(bp_doc, '_height', 600))
    scale = min(1280 / doc_w, 720 / doc_h)
    offset_x = (1280 - doc_w * scale) / 2
    offset_y = (720 - doc_h * scale) / 2

    num_slides = bp_doc._global_counter.get('slide', -1) + 1
    svg_files = []
    all_plain_texts = []

    for islide in range(num_slides):
        slide_id = f"slide_{islide}"
        slide_obj = bp_doc._slides[slide_id]
        slide_obj.newrender()
        
        current_plain_text_mods = {}
        current_plain_text_ids = set()

        if hasattr(slide_obj, 'title_element') and slide_obj.title_element:
            t = slide_obj.title_element
            if not is_math(t.content):
                current_plain_text_mods[t.id] = t
                current_plain_text_ids.add(t.id)

        for mid in slide_obj.element_keys:
            m = slide_obj.contents[mid]
            if m.type == 'text' and not is_math(m.content):
                current_plain_text_mods[m.id] = m
                current_plain_text_ids.add(m.id)
        
        all_plain_texts.append(current_plain_text_mods)

        glyphs_svg = ""
        if 'glyphs' in bp_doc._global_store:
            glyphs_svg = '<defs>' + ''.join([g['svg'] for g in bp_doc._global_store['glyphs'].values()]) + '</defs>'
        def_svg = '<defs>' + ''.join(slide_obj.svgdefout) + '</defs>'
        content = slide_obj.svglayers.get(0, "")
        full_svg_str = slide_obj.svgheader + glyphs_svg + def_svg + content + slide_obj.svgfooter

        # Save Native Version (Full vector including text as paths - Ground Truth)
        native_svg_path = debug_dir / f"native_p{islide+1:02d}.svg"
        with open(native_svg_path, "w", encoding="utf-8") as f:
            f.write(full_svg_str)

        # Create Wrapped Version (Headless for PPTX shapes)
        root = ET.fromstring(full_svg_str.encode('utf-8'))
        flat_root = absolute_flatten(root, plain_text_ids=current_plain_text_ids, target_w=1280, target_h=720, doc_w=doc_w, doc_h=doc_h)
        flat_svg_str = ET.tostring(flat_root, encoding='unicode')
        
        svg_path = debug_dir / f"wrapped_p{islide+1:02d}.svg"
        with open(svg_path, "w", encoding="utf-8") as f:
            f.write(flat_svg_str)
        svg_files.append(svg_path)

    create_pptx_with_native_svg(svg_files=svg_files, output_path=Path(output_pptx), use_native_shapes=True, verbose=False)
    
    prs = Presentation(output_pptx)
    for islide in range(num_slides):
        slide_obj = bp_doc._slides[f"slide_{islide}"]
        pptx_slide = prs.slides[islide]
        plain_mods = all_plain_texts[islide]
        
        for mid, m in plain_mods.items():
            bx = m.positionner.x['final']
            by = m.positionner.y['final']
            bw = m.positionner.width.value
            bh = m.height.value if m.height.value else m.size * 1.5
            
            px, py = bx * scale + offset_x, by * scale + offset_y
            pw, ph = bw * scale, bh * scale
            
            txBox = pptx_slide.shapes.add_textbox(Inches(px/96), Inches(py/96), Inches(pw/96), Inches(ph/96))
            tf = txBox.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.text = clean_latex_for_pptx(m.content)
            
            if hasattr(m, 'size'): p.font.size = Pt(m.size * scale * 0.75)
            if hasattr(m, 'color'):
                try:
                    c_hex = mcolors.to_hex(m.color).lstrip('#')
                    p.font.color.rgb = RGBColor(int(c_hex[0:2], 16), int(c_hex[2:4], 16), int(c_hex[4:6], 16))
                except: pass
            if hasattr(m, 'align'):
                if 'center' in str(m.align): p.alignment = PP_ALIGN.CENTER
                elif 'right' in str(m.align): p.alignment = PP_ALIGN.RIGHT

    prs.save(output_pptx)
    print(f"\n✅ Professional Presentation (Cleaned Text & Native Math) saved to {output_pptx}")

if __name__ == "__main__":
    pass
